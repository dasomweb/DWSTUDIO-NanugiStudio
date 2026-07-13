"""입력 파일 파싱 (DW-ListPilot `services/excel_parser.py` · `image_pipeline.process_image` 이식).

엑셀·CSV 는 표 그대로 읽고, 이미지는 Shopify 권장 크기로 줄여 WebP 로 바꾼다.
PDF 는 페이지를 이미지로 렌더해 Gemini 비전에 넘긴다 (gemini.py 가 이어받는다).
"""

from __future__ import annotations

import csv
import io

MAX_IMAGE_SIDE = 2048  # Shopify 권장 상한


def parse_csv(file_bytes: bytes) -> list[dict]:
    text = file_bytes.decode("utf-8-sig")
    return [dict(row) for row in csv.DictReader(io.StringIO(text))]


def parse_excel(file_bytes: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(h).strip().lower() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    out: list[dict] = []
    for row in rows[1:]:
        if not any(row):
            continue
        out.append({headers[i]: v for i, v in enumerate(row) if i < len(headers)})
    return out


def pdf_page_images(file_bytes: bytes, dpi: int = 300) -> list[bytes]:
    """PDF 페이지를 PNG 로 렌더한다. 텍스트 추출이 아니라 이미지로 넘기는 이유는,
    도매 인보이스·카탈로그가 표·이미지 혼합이라 비전 모델이 더 잘 읽기 때문이다."""
    import fitz  # PyMuPDF

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    scale = dpi / 72
    pages = [doc[i].get_pixmap(matrix=fitz.Matrix(scale, scale)).tobytes("png") for i in range(len(doc))]
    doc.close()
    return pages


def process_image(image_bytes: bytes, max_size: int = MAX_IMAGE_SIDE) -> bytes:
    """긴 변이 max_size 를 넘으면 줄이고, WebP 로 변환하며 메타데이터를 떨군다."""
    from PIL import Image

    img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    if max(img.size) > max_size:
        img.thumbnail((max_size, max_size), Image.LANCZOS)

    out = io.BytesIO()
    img.save(out, format="WEBP", quality=85, optimize=True)
    return out.getvalue()


# 도매처마다 컬럼명이 다르다. 흔한 표기를 내부 필드로 흡수한다.
COLUMN_ALIASES: dict[str, list[str]] = {
    "title": ["title", "product name", "name", "item", "product", "item name"],
    "sku": ["sku", "item code", "product code", "code", "style number", "style #"],
    "barcode": ["barcode", "upc", "ean", "gtin"],
    "price": ["price", "unit price", "retail price", "wholesale price", "cost"],
    "vendor": ["vendor", "brand", "manufacturer", "supplier"],
    "quantity": ["qty", "quantity", "stock", "inventory", "units"],
    "description": ["description", "desc", "details", "body"],
    "weight": ["weight", "weight_kg", "weight (kg)", "weight (g)"],
    "size": ["size", "sizes"],
    "color": ["color", "colour", "colors"],
}


def auto_map_columns(headers: list[str]) -> dict[str, str]:
    """표 헤더 → 내부 필드명. 못 찾은 필드는 빠진다 (호출부가 없는 값으로 다룬다)."""
    normalized = {h.strip().lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[field] = normalized[alias]
                break
    return mapping


def row_to_extracted(row: dict, mapping: dict[str, str]) -> dict:
    """표 한 줄 → Gemini 추출 결과와 같은 모양. 이후 파이프라인이 둘을 구분하지 않게 한다."""

    def get(field: str):
        col = mapping.get(field)
        return row.get(col) if col else None

    def num(field: str):
        v = get(field)
        if v in (None, ""):
            return None
        try:
            return float(str(v).replace(",", "").replace("$", "").strip())
        except ValueError:
            return None

    options = []
    for axis in ("size", "color"):
        raw = get(axis)
        if raw:
            values = [v.strip() for v in str(raw).replace("/", ",").split(",") if v.strip()]
            if values:
                options.append({"name": axis.capitalize(), "values": values})

    return {
        "title": str(get("title") or "Untitled Product"),
        "vendor": str(get("vendor")) if get("vendor") else None,
        "product_type": None,
        "description": str(get("description") or ""),
        "price": num("price"),
        "sku": str(get("sku")) if get("sku") else None,
        "barcode": str(get("barcode")) if get("barcode") else None,
        "weight_kg": num("weight"),
        "quantity": int(num("quantity") or 0),
        "options": options,
        "tags_hints": {k: str(get(k)) for k in ("color",) if get(k)},
        "confidence": 1.0,  # 사람이 만든 표다. 추측이 아니다.
    }
